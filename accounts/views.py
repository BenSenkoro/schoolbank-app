from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth import authenticate, login, logout
from django.contrib import messages
from .models import CustomUser
from banking.models import BankAccount
from django.db.models import Q
import csv
import io
from django.http import HttpResponse
from django.contrib.auth.decorators import login_required
from .utils import log_activity, paginate

def role_required(*roles):
    def decorator(view_func):
        def wrapper(request, *args, **kwargs):
            if not request.user.is_authenticated:
                return redirect('login')
            if request.user.role not in roles:
                messages.error(request, 'You do not have permission to access that page.')
                return redirect('dashboard')
            return view_func(request, *args, **kwargs)
        return wrapper
    return decorator


def login_view(request):
    if request.user.is_authenticated:
        return redirect('dashboard')

    if request.method == 'POST':
        username = request.POST.get('username')
        password = request.POST.get('password')
        user = authenticate(request, username=username, password=password)
        if user is not None:
            login(request, user)
            log_activity(
                user=user,
                action='login',
                description=f'{user.get_full_name() or user.username} logged in.',
                request=request
            )
            return redirect('dashboard')
        else:
            messages.error(request, 'Invalid username or password.')

    return render(request, 'accounts/login.html')


def logout_view(request):
    if request.user.is_authenticated:
        log_activity(
            user=request.user,
            action='logout',
            description=f'{request.user.get_full_name() or request.user.username} logged out.',
            request=request
        )
    logout(request)
    return redirect('login')


@role_required('admin')
def user_list(request):
    query = request.GET.get('q', '').strip()
    role_filter = request.GET.get('role', '').strip()

    users = CustomUser.objects.all().order_by('role', 'username')

    if query:
        users = users.filter(
            __import__('django.db.models', fromlist=['Q']).Q(username__icontains=query) |
            __import__('django.db.models', fromlist=['Q']).Q(first_name__icontains=query) |
            __import__('django.db.models', fromlist=['Q']).Q(last_name__icontains=query)
        )

    if role_filter:
        users = users.filter(role=role_filter)

    summary = {
        'total': CustomUser.objects.count(),
        'admins': CustomUser.objects.filter(role='admin').count(),
        'managers': CustomUser.objects.filter(role='manager').count(),
        'loan_managers': CustomUser.objects.filter(role='loan_manager').count(),
        'students': CustomUser.objects.filter(role='student').count(),
        'teachers': CustomUser.objects.filter(role='teacher').count(),
        'active': CustomUser.objects.filter(is_active=True).count(),
    }

    page_obj = paginate(request, users)
    return render(request, 'accounts/user_list.html', {
        'users': page_obj,
        'page_obj': page_obj,
        'query': query,
        'role_filter': role_filter,
        'summary': summary,
    })


@role_required('admin')
def create_user(request):
    if request.method == 'POST':
        username = request.POST.get('username', '').strip()
        first_name = request.POST.get('first_name', '').strip()
        last_name = request.POST.get('last_name', '').strip()
        email = request.POST.get('email', '').strip()
        role = request.POST.get('role')
        password = request.POST.get('password')
        confirm_password = request.POST.get('confirm_password')

        if not username or not password or not role:
            messages.error(request, 'Username, password and role are required.')
        elif password != confirm_password:
            messages.error(request, 'Passwords do not match.')
        elif CustomUser.objects.filter(username=username).exists():
            messages.error(request, f'Username "{username}" is already taken.')
        else:
            try:
                user = CustomUser.objects.create_user(
                    username=username,
                    first_name=first_name,
                    last_name=last_name,
                    email=email,
                    password=password,
                    role=role
                )
                log_activity(
                user=request.user,
                action='user_created',
                description=f'User "{username}" created with role {user.get_role_display()}.',
                request=request
)
                messages.success(request, f'User "{username}" created successfully as {user.get_role_display()}.')
                return redirect('user_list')
            except Exception as e:
                messages.error(request, f'Error creating user: {e}')

    return render(request, 'accounts/create_user.html')


@role_required('admin')
def edit_user(request, pk):
    user = get_object_or_404(CustomUser, pk=pk)

    if request.method == 'POST':
        user.first_name = request.POST.get('first_name', '').strip()
        user.last_name = request.POST.get('last_name', '').strip()
        user.email = request.POST.get('email', '').strip()
        user.role = request.POST.get('role')

        new_password = request.POST.get('new_password', '').strip()
        confirm_password = request.POST.get('confirm_password', '').strip()

        if new_password:
            if new_password != confirm_password:
                messages.error(request, 'Passwords do not match.')
                return render(request, 'accounts/edit_user.html', {'edited_user': user})
            user.set_password(new_password)

        user.save()
        messages.success(request, f'User "{user.username}" updated successfully.')
        return redirect('user_list')
        log_activity(
        user=request.user,
        action='user_edited',
        description=f'User "{user.username}" details updated.',
        request=request)

    return render(request, 'accounts/edit_user.html', {'edited_user': user})


@role_required('admin')
def toggle_user_status(request, pk):
    user = get_object_or_404(CustomUser, pk=pk)

    if user == request.user:
        messages.error(request, 'You cannot deactivate your own account.')
        return redirect('user_list')

    user.is_active = not user.is_active
    status = 'activated' if user.is_active else 'deactivated'
    user.save()
    log_activity(
        user=request.user,
        action='user_toggled',
        description=f'User "{user.username}" was {status}.',
        request=request
    )
    messages.success(request, f'User "{user.username}" has been {status}.')
    return redirect('user_list')
@role_required('admin')
def bulk_import(request):
    if request.method == 'POST':
        uploaded_file = request.FILES.get('file')
        create_bank_accounts = request.POST.get('create_bank_accounts') == 'yes'

        if not uploaded_file:
            messages.error(request, 'Please select a file to upload.')
            return redirect('bulk_import')

        filename = uploaded_file.name.lower()

        if not (filename.endswith('.csv') or filename.endswith('.xlsx')):
            messages.error(request, 'Only .csv and .xlsx files are supported.')
            return redirect('bulk_import')

        rows = []

        try:
            if filename.endswith('.csv'):
                decoded = uploaded_file.read().decode('utf-8-sig')
                reader = csv.DictReader(io.StringIO(decoded))
                for row in reader:
                    rows.append(row)

            elif filename.endswith('.xlsx'):
                import openpyxl
                wb = openpyxl.load_workbook(uploaded_file)
                ws = wb.active
                headers = [cell.value for cell in ws[1]]
                for row in ws.iter_rows(min_row=2, values_only=True):
                    rows.append(dict(zip(headers, row)))

        except Exception as e:
            messages.error(request, f'Error reading file: {e}')
            return redirect('bulk_import')

        # Process rows
        created = []
        skipped = []
        failed = []

        for i, row in enumerate(rows, start=2):
            first_name = str(row.get('first_name') or '').strip()
            last_name = str(row.get('last_name') or '').strip()
            username = str(row.get('username') or '').strip()
            email = str(row.get('email') or '').strip()
            role = str(row.get('role') or '').strip().lower()
            password = str(row.get('password') or '').strip()

            # Parse opening balance
            raw_balance = row.get('opening_balance')
            try:
                opening_balance = float(str(raw_balance).strip()) if raw_balance not in [None, '', 'None'] else 0.0
            except (ValueError, TypeError):
                opening_balance = 0.0

            # Validate required fields
            if not username or not password or not role:
                failed.append({
                    'row': i,
                    'username': username or '—',
                    'reason': 'Missing username, password or role'
                })
                continue

            # Validate role
            valid_roles = ['admin', 'manager', 'loan_manager', 'teller', 'student', 'teacher']
            if role not in valid_roles:
                failed.append({
                    'row': i,
                    'username': username,
                    'reason': f'Invalid role "{role}". Must be one of: {", ".join(valid_roles)}'
                })
                continue

            # Check if username already exists
            if CustomUser.objects.filter(username=username).exists():
                skipped.append({
                    'row': i,
                    'username': username,
                    'reason': 'Username already exists'
                })
                continue

            try:
                # Create user
                user = CustomUser.objects.create_user(
                    username=username,
                    first_name=first_name,
                    last_name=last_name,
                    email=email,
                    password=password,
                    role=role
                )

                bank_account_created = False
                account_number = '—'

                # Create bank account if requested and role qualifies
                if create_bank_accounts and role in ['student', 'teacher']:
                    from banking.models import BankAccount, Transaction
                    from django.utils import timezone

                    account = BankAccount.objects.create(
                        user=user,
                        account_type=role,
                        balance=opening_balance
                    )
                    bank_account_created = True
                    account_number = account.account_number

                    # Record opening balance as a transaction if greater than 0
                    if opening_balance > 0:
                        Transaction.objects.create(
                            account=account,
                            transaction_type='deposit',
                            amount=opening_balance,
                            performed_by=request.user,
                            notes='Opening balance migrated from previous system',
                            timestamp=timezone.now()
                        )

                created.append({
                    'username': username,
                    'name': f'{first_name} {last_name}'.strip(),
                    'role': user.get_role_display(),
                    'opening_balance': opening_balance if role in ['student', 'teacher'] else '—',
                    'account_number': account_number if bank_account_created else '—',
                })

            except Exception as e:
                failed.append({
                    'row': i,
                    'username': username,
                    'reason': str(e)
                })

        log_activity(
            user=request.user,
            action='bulk_import',
            description=f'Bulk import: {len(created)} created, {len(skipped)} skipped, {len(failed)} failed.',
            request=request
        )

        return render(request, 'accounts/import_results.html', {
            'created': created,
            'skipped': skipped,
            'failed': failed,
            'create_bank_accounts': create_bank_accounts,
        })

    return render(request, 'accounts/bulk_import.html')


@role_required('admin')
def download_template(request):
    file_type = request.GET.get('type', 'csv')

    if file_type == 'xlsx':
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'Users Import'

        # Headers
        headers = ['first_name', 'last_name', 'username', 'email', 'role', 'password', 'opening_balance']
        header_fill = PatternFill(start_color='0077B6', end_color='0077B6', fill_type='solid')
        header_font = Font(bold=True, color='FFFFFF')

        for col, header in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center')
            ws.column_dimensions[cell.column_letter].width = 20

        # Sample rows
        sample_rows = [
            ['John', 'Doe', 'john_doe', 'john@mbonea.ac.tz', 'student', 'pass1234', 50000],
            ['Jane', 'Smith', 'jane_smith', 'jane@mbonea.ac.tz', 'teacher', 'pass1234', 120000],
            ['Ali', 'Hassan', 'ali_hassan', 'ali@mbonea.ac.tz', 'student', 'pass1234', 0],
            ['', '', 'manager1', '', 'manager', 'pass1234', ''],
        ]

        for row_data in sample_rows:
            ws.append(row_data)

        # Notes sheet
        ws2 = wb.create_sheet('Notes')
        ws2['A1'] = 'IMPORT NOTES'
        ws2['A1'].font = Font(bold=True, size=13)
        notes = [
            ('first_name', 'Optional — first name of the user'),
            ('last_name', 'Optional — last name of the user'),
            ('username', 'REQUIRED — must be unique, no spaces'),
            ('email', 'Optional — email address'),
            ('role', 'REQUIRED — must be one of: admin, manager, loan_manager, teller, student, teacher'),
            ('password', 'REQUIRED — initial password for the user'),
            ('opening_balance', 'Optional — amount already in account from previous system. Leave blank or 0 if none. Only applies to students and teachers.'),
        ]
        for i, (field, note) in enumerate(notes, start=3):
            ws2.cell(row=i, column=1, value=field).font = Font(bold=True)
            ws2.cell(row=i, column=2, value=note)
            ws2.column_dimensions['A'].width = 20
            ws2.column_dimensions['B'].width = 65

        response = HttpResponse(
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = 'attachment; filename="mbonea_users_template.xlsx"'
        wb.save(response)
        return response

    else:
        # CSV template
        response = HttpResponse(content_type='text/csv')
        response['Content-Disposition'] = 'attachment; filename="mbonea_users_template.csv"'
        writer = csv.writer(response)
        writer.writerow(['first_name', 'last_name', 'username', 'email', 'role', 'password', 'opening_balance'])
        writer.writerow(['John', 'Doe', 'john_doe', 'john@mbonea.ac.tz', 'student', 'pass1234', 50000])
        writer.writerow(['Jane', 'Smith', 'jane_smith', 'jane@mbonea.ac.tz', 'teacher', 'pass1234', 120000])
        writer.writerow(['Ali', 'Hassan', 'ali_hassan', 'ali@mbonea.ac.tz', 'student', 'pass1234', 0])
        writer.writerow(['', '', 'manager1', '', 'manager', 'pass1234', ''])
        return response
@role_required('admin')
def activity_log(request):
    from .models import ActivityLog
    query = request.GET.get('q', '').strip()
    action_filter = request.GET.get('action', '').strip()

    logs = ActivityLog.objects.select_related('user').all()

    if query:
        logs = logs.filter(
            Q(user__username__icontains=query) |
            Q(user__first_name__icontains=query) |
            Q(description__icontains=query)
        )

    if action_filter:
        logs = logs.filter(action=action_filter)

    page_obj = paginate(request, logs)
    return render(request, 'accounts/activity_log.html', {
        'logs': page_obj,
        'page_obj': page_obj,
        'query': query,
        'action_filter': action_filter,
    })
from django.contrib.auth import update_session_auth_hash

@login_required
def change_password(request):
    if request.method == 'POST':
        current_password = request.POST.get('current_password')
        new_password = request.POST.get('new_password')
        confirm_password = request.POST.get('confirm_password')

        if not request.user.check_password(current_password):
            messages.error(request, 'Current password is incorrect.')
        elif len(new_password) < 6:
            messages.error(request, 'New password must be at least 6 characters.')
        elif new_password != confirm_password:
            messages.error(request, 'New passwords do not match.')
        else:
            request.user.set_password(new_password)
            request.user.save()
            update_session_auth_hash(request, request.user)
            log_activity(
                user=request.user,
                action='user_edited',
                description=f'{request.user.username} changed their password.',
                request=request
            )
            messages.success(request, 'Password changed successfully.')
            return redirect('change_password')

    return render(request, 'accounts/change_password.html')